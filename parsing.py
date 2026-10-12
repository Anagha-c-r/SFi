"""Read uploaded CSV / Excel files into a list of header->value dicts."""
from __future__ import annotations

import csv
import io
from dataclasses import dataclass


class ParseError(ValueError):
    pass


@dataclass
class ParsedFile:
    headers: list[str]
    rows: list[dict[str, str]]


MAX_ROWS = 200_000


def _decode(data: bytes) -> str:
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    raise ParseError("Could not read the file's text encoding. Save it as UTF-8 CSV and try again.")


def _cell(value) -> str:
    if value is None:
        return ""
    if hasattr(value, "strftime"):
        # openpyxl returns datetime for date cells; keep ISO so the validator can parse it
        return value.strftime("%Y-%m-%d")
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def parse_csv(data: bytes) -> ParsedFile:
    text = _decode(data)
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    reader = csv.reader(io.StringIO(text), dialect)
    try:
        headers = [h.strip() for h in next(reader)]
    except StopIteration:
        raise ParseError("The file is empty.")
    rows = []
    for values in reader:
        if not any(v.strip() for v in values):
            continue
        rows.append({h: (values[i].strip() if i < len(values) else "") for i, h in enumerate(headers)})
        if len(rows) > MAX_ROWS:
            raise ParseError(f"The file has more than {MAX_ROWS:,} rows. Split it into smaller files.")
    return ParsedFile(headers=headers, rows=rows)


def parse_xlsx(data: bytes) -> ParsedFile:
    from openpyxl import load_workbook

    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # openpyxl raises many different exception types
        raise ParseError(f"Could not open the Excel file: {exc}")
    ws = wb.worksheets[0]
    it = ws.iter_rows(values_only=True)
    try:
        headers = [_cell(h) for h in next(it)]
    except StopIteration:
        raise ParseError("The first sheet is empty.")
    while headers and not headers[-1]:
        headers.pop()
    rows = []
    for values in it:
        cells = [_cell(v) for v in values[: len(headers)]]
        if not any(cells):
            continue
        rows.append({h: (cells[i] if i < len(cells) else "") for i, h in enumerate(headers)})
        if len(rows) > MAX_ROWS:
            raise ParseError(f"The file has more than {MAX_ROWS:,} rows. Split it into smaller files.")
    wb.close()
    return ParsedFile(headers=headers, rows=rows)


def parse_upload(filename: str, data: bytes) -> ParsedFile:
    name = (filename or "").lower()
    if name.endswith((".xlsx", ".xlsm")):
        return parse_xlsx(data)
    if name.endswith((".csv", ".txt")):
        return parse_csv(data)
    raise ParseError("Upload a .csv or .xlsx file.")
